#!/usr/bin/env python3
"""Final Prototype-Phase deck = the submitted Round 1 deck's slides 1-8 (story, unchanged) + this deck's slides 9-25.
    python deck/build_deck.py && python deck/assemble.py
Round 1 pages come from ../../../submission/deck/Saqqa_Pitch_Deck.pdf, the file that was actually uploaded."""
import pathlib, re
import fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
HERE = pathlib.Path(__file__).parent
R1 = HERE.parent.parent.parent / "submission" / "deck" / "Saqqa_Pitch_Deck.pdf"
order = re.findall(r'^slide\("([^"]+)"', (HERE / "build_deck.py").read_text(), re.M)
tmp = HERE / "png" / "r1"; tmp.mkdir(parents=True, exist_ok=True)
doc = fitz.open(R1); seq = []
for i in range(8):
    p = tmp / f"r1_{i+1:02d}.png"; doc[i].get_pixmap(matrix=fitz.Matrix(1920 / doc[i].rect.width, 1080 / doc[i].rect.height)).save(p); seq.append(p)
seq += [HERE / "png" / f"{n}.png" for n in order[8:]]
imgs = [Image.open(p).convert("RGB") for p in seq]
out = HERE / "out"; pdf = out / "Saqqa_Prototype_Deck.pdf"
imgs[0].save(pdf, save_all=True, append_images=imgs[1:], resolution=144.0, quality=92)
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
for p in seq:
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(p), 0, 0, prs.slide_width, prs.slide_height)
prs.save(str(out / "Saqqa_Prototype_Deck.pptx"))
print(len(seq), "slides ->", pdf)
